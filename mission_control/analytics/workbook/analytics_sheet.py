"""Study Intelligence workbook sheet."""

from __future__ import annotations

from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font

from mission_control.analytics.models import StudyIntelligenceReport
from mission_control.workbook.table_sheet import TableSheet


class StudyIntelligenceSheet(TableSheet):
    """Render analytics outputs without calculating business logic."""

    title = "Study Intelligence"

    def build(self) -> None:
        """Build the Study Intelligence worksheet."""
        self.prepare_sheet(subtitle="Analytics, recommendations, and readiness")
        report = self._report()

        stats_end_row, _ = self._write_overall_statistics(report)
        velocity_end_row, _ = self._write_learning_velocity(report, stats_end_row + 3)
        weak_end_row, _ = self._write_weak_topics(report, velocity_end_row + 3)
        rec_end_row, _ = self._write_recommendations(report, weak_end_row + 3)
        readiness_end_row, _ = self._write_readiness(report, rec_end_row + 3)
        distribution_end_row, _ = self._write_distribution(
            report,
            readiness_end_row + 3,
        )
        trend_end_row, _ = self._write_hours_trend(report, distribution_end_row + 3)

        self.freeze_header_row(4)
        self.auto_size_columns()
        self._add_charts(
            stats_start_row=5,
            distribution_start_row=(
                distribution_end_row - len(report.study_distribution) + 1
            ),
            distribution_end_row=distribution_end_row,
            trend_start_row=trend_end_row - len(report.study_hours_trend) + 1,
            trend_end_row=trend_end_row,
        )
        self._apply_conditional_formatting(
            start_row=weak_end_row - len(report.weak_topics) + 1,
            end_row=weak_end_row,
        )
        self._apply_page_setup()

    def _report(self) -> StudyIntelligenceReport:
        if self.context is None:
            raise RuntimeError("Study Intelligence requires application context.")
        return self.context.study_intelligence

    def _write_section_title(self, row: int, title: str) -> None:
        cell = self.worksheet.cell(row=row, column=1, value=title)
        cell.font = Font(
            name=self.theme.header_font,
            size=self.theme.header_size,
            bold=True,
            color=self.theme.primary_dark,
        )

    def _write_overall_statistics(
        self,
        report: StudyIntelligenceReport,
    ) -> tuple[int, int]:
        self._write_section_title(4, "Overall Statistics")
        return self.write_table(
            ["Metric", "Value"],
            [
                ("Total Tasks", report.statistics.total_tasks),
                ("Completed Tasks", report.statistics.completed_tasks),
                ("Pending Tasks", report.statistics.pending_tasks),
                ("Overdue Tasks", report.statistics.overdue_tasks),
                ("Total Hours", report.statistics.total_hours),
                ("Completed Hours", report.statistics.completed_hours),
                ("Completion %", report.statistics.completion_percentage),
                (
                    "Revision Completion %",
                    report.statistics.revision_completion_percentage,
                ),
            ],
            start_row=5,
        )

    def _write_learning_velocity(
        self,
        report: StudyIntelligenceReport,
        start_row: int,
    ) -> tuple[int, int]:
        self._write_section_title(start_row, "Learning Velocity")
        projected = report.velocity.projected_completion_date
        return self.write_table(
            ["Metric", "Value"],
            [
                ("Study Hours / Day", report.velocity.study_hours_per_day),
                ("Questions / Day", report.velocity.questions_per_day),
                ("Questions / Hour", report.velocity.questions_per_hour),
                (
                    "Average Completion Rate",
                    report.velocity.average_completion_rate,
                ),
                (
                    "Projected Completion Date",
                    projected.isoformat() if projected else "Insufficient data",
                ),
            ],
            start_row=start_row + 1,
        )

    def _write_weak_topics(
        self,
        report: StudyIntelligenceReport,
        start_row: int,
    ) -> tuple[int, int]:
        self._write_section_title(start_row, "Weak Topics")
        rows = [
            (
                topic.topic,
                topic.attempts,
                topic.accuracy,
                topic.confidence,
                topic.priority,
            )
            for topic in report.weak_topics
        ]
        return self.write_table(
            ["Topic", "Attempts", "Accuracy", "Confidence", "Priority"],
            rows,
            start_row=start_row + 1,
        )

    def _write_recommendations(
        self,
        report: StudyIntelligenceReport,
        start_row: int,
    ) -> tuple[int, int]:
        self._write_section_title(start_row, "Recommendations")
        rows = [
            (
                recommendation.title,
                recommendation.description,
                recommendation.priority,
                recommendation.category,
                recommendation.generated_date,
            )
            for recommendation in report.recommendations
        ]
        return self.write_table(
            ["Title", "Description", "Priority", "Category", "Generated Date"],
            rows,
            start_row=start_row + 1,
        )

    def _write_readiness(
        self,
        report: StudyIntelligenceReport,
        start_row: int,
    ) -> tuple[int, int]:
        self._write_section_title(start_row, "Readiness")
        return self.write_table(
            ["Metric", "Value"],
            [
                ("Score", report.readiness.score),
                ("Completion Weight", report.readiness.completion_weight),
                ("Revision Weight", report.readiness.revision_weight),
                ("Consistency Weight", report.readiness.consistency_weight),
                ("Mock Weight", report.readiness.mock_weight),
                ("Confidence Weight", report.readiness.confidence_weight),
            ],
            start_row=start_row + 1,
        )

    def _write_distribution(
        self,
        report: StudyIntelligenceReport,
        start_row: int,
    ) -> tuple[int, int]:
        self._write_section_title(start_row, "Study Distribution")
        rows = [
            (item.subject, item.task_count, item.planned_hours)
            for item in report.study_distribution
        ]
        return self.write_table(
            ["Subject", "Tasks", "Planned Hours"],
            rows,
            start_row=start_row + 1,
        )

    def _write_hours_trend(
        self,
        report: StudyIntelligenceReport,
        start_row: int,
    ) -> tuple[int, int]:
        self._write_section_title(start_row, "Study Hours Trend")
        rows = [
            (item.day, item.planned_hours, item.completed_hours)
            for item in report.study_hours_trend
        ]
        return self.write_table(
            ["Day", "Planned Hours", "Completed Hours"],
            rows,
            start_row=start_row + 1,
        )

    def _add_charts(
        self,
        *,
        stats_start_row: int,
        distribution_start_row: int,
        distribution_end_row: int,
        trend_start_row: int,
        trend_end_row: int,
    ) -> None:
        self._add_completed_pending_bar_chart(stats_start_row)
        self._add_distribution_pie_chart(
            distribution_start_row,
            distribution_end_row,
        )
        self._add_hours_trend_line_chart(trend_start_row, trend_end_row)

    def _add_completed_pending_bar_chart(self, stats_start_row: int) -> None:
        chart = BarChart()
        chart.title = "Completed vs Pending Tasks"
        chart.y_axis.title = "Tasks"
        data = Reference(
            self.worksheet,
            min_col=2,
            min_row=stats_start_row + 2,
            max_row=stats_start_row + 3,
        )
        categories = Reference(
            self.worksheet,
            min_col=1,
            min_row=stats_start_row + 2,
            max_row=stats_start_row + 3,
        )
        chart.add_data(data)
        chart.set_categories(categories)
        self.worksheet.add_chart(chart, "H5")

    def _add_distribution_pie_chart(
        self,
        start_row: int,
        end_row: int,
    ) -> None:
        if end_row < start_row:
            return
        chart = PieChart()
        chart.title = "Study Distribution"
        data = Reference(
            self.worksheet,
            min_col=3,
            min_row=start_row,
            max_row=end_row,
        )
        labels = Reference(
            self.worksheet,
            min_col=1,
            min_row=start_row,
            max_row=end_row,
        )
        chart.add_data(data)
        chart.set_categories(labels)
        self.worksheet.add_chart(chart, "H20")

    def _add_hours_trend_line_chart(
        self,
        start_row: int,
        end_row: int,
    ) -> None:
        if end_row < start_row:
            return
        chart = LineChart()
        chart.title = "Study Hours Trend"
        chart.y_axis.title = "Hours"
        data = Reference(
            self.worksheet,
            min_col=2,
            max_col=3,
            min_row=start_row - 1,
            max_row=end_row,
        )
        categories = Reference(
            self.worksheet,
            min_col=1,
            min_row=start_row,
            max_row=end_row,
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        self.worksheet.add_chart(chart, "H35")

    def _apply_conditional_formatting(self, *, start_row: int, end_row: int) -> None:
        if end_row < start_row:
            return
        self.worksheet.conditional_formatting.add(
            f"C{start_row}:C{end_row}",
            ColorScaleRule(
                start_type="num",
                start_value=0,
                start_color="F8696B",
                mid_type="num",
                mid_value=70,
                mid_color="FFEB84",
                end_type="num",
                end_value=100,
                end_color="63BE7B",
            ),
        )
